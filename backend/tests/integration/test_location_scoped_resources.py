from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

import httpx
import pdfplumber
from openpyxl import load_workbook
from sqlalchemy import delete, select, text

from assetguard.app import app
from assetguard.infrastructure.database import get_session_factory
from assetguard.modules.assets.models import (
    AssetRecord, BuildingRecord, FloorRecord, OrganizationRecord,
    RoomInspectionItemRecord, RoomInspectionRecord, RoomRecord,
)
from assetguard.modules.identity.auth import hash_password
from assetguard.modules.identity.models import AuthSessionRecord, LocationAccessRecord, UserRecord
from assetguard.modules.incidents.models import AssetHistoryEntryRecord
from assetguard.modules.vision.models import VisionBaselineRecord, VisionRoomRecord, VisionScanRecord


def test_location_scopes_cover_exports_and_vision_resources() -> None:
    asyncio.run(_exercise_location_scope())


async def _exercise_location_scope() -> None:
    now = datetime.now(UTC)
    factory = get_session_factory()
    building_ids = []
    floor_ids = []
    room_ids = []
    with factory() as session:
        organization = OrganizationRecord(name=f"Scope School {uuid4().hex[:8]}", created_at=now)
        session.add(organization)
        session.flush()

        def add_room(building_name: str, room_name: str) -> RoomRecord:
            building = BuildingRecord(organization_id=organization.id, name=building_name, created_at=now)
            session.add(building)
            session.flush()
            building_ids.append(building.id)
            floor = FloorRecord(building_id=building.id, name="1", created_at=now)
            session.add(floor)
            session.flush()
            floor_ids.append(floor.id)
            room = RoomRecord(floor_id=floor.id, name=room_name, created_at=now)
            session.add(room)
            session.flush()
            room_ids.append(room.id)
            return room

        visible_room = add_room("Scope корпус A", "101")
        hidden_room = add_room("Scope корпус B", "201")
        visible_asset = AssetRecord(
                organization_id=organization.id, room_id=visible_room.id, inventory_number="SCOPE-VISIBLE",
                name="Visible school device", asset_type="Desktop", category="IT", tracking_mode="INDIVIDUAL",
                quantity=1, unit="шт.", status="ACTIVE", building="Scope корпус A", floor="1", room="101",
                created_at=now, updated_at=now,
            )
        hidden_asset = AssetRecord(
                organization_id=organization.id, room_id=hidden_room.id, inventory_number="SCOPE-HIDDEN",
                name="Hidden school device", asset_type="Desktop", category="IT", tracking_mode="INDIVIDUAL",
                quantity=1, unit="шт.", status="ACTIVE", building="Scope корпус B", floor="1", room="201",
                created_at=now, updated_at=now,
            )
        session.add_all([visible_asset, hidden_asset])
        visible_vision_room = VisionRoomRecord(
            name="Scope корпус A / 1 / 101", organization_id=organization.id,
            location_room_id=visible_room.id, created_at=now,
        )
        hidden_vision_room = VisionRoomRecord(
            name="Scope корпус B / 1 / 201", organization_id=organization.id,
            location_room_id=hidden_room.id, created_at=now,
        )
        legacy_vision_room = VisionRoomRecord(name="Legacy unlinked room", organization_id=organization.id, created_at=now)
        session.add_all([visible_vision_room, hidden_vision_room, legacy_vision_room])
        session.flush()
        hidden_scan = VisionScanRecord(
            room_id=hidden_vision_room.id, created_at=now, status="OK",
            original_image_path="unused.jpg", annotated_image_path="unused.jpg",
            counts={}, comparison={}, model_id="test", confidence_threshold=0.5,
        )
        mismatched_legacy_scan = VisionScanRecord(
            room_id=visible_vision_room.id, asset_id=hidden_asset.id, created_at=now, status="OK",
            original_image_path="unused.jpg", annotated_image_path="unused.jpg",
            counts={}, comparison={}, model_id="test", confidence_threshold=0.5,
        )
        session.add_all([hidden_scan, mismatched_legacy_scan])
        user = UserRecord(
            username=f"scope-viewer-{uuid4().hex[:8]}", password_hash=hash_password("scope-password-123"),
            role="LOCATION_MANAGER", organization_id=organization.id, is_active=True, created_at=now,
        )
        session.add(user)
        session.flush()
        created_user_id = user.id
        session.add(LocationAccessRecord(
            user_id=user.id, scope_type="ROOM", scope_id=visible_room.id,
            permission="VIEWER", created_at=now,
        ))
        session.commit()
        visible_room_id, hidden_room_id = visible_room.id, hidden_vision_room.id
        visible_asset_id, hidden_scan_id = visible_asset.id, hidden_scan.id
        username = user.username

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        login = await client.post("/auth/login", json={"username": username, "password": "scope-password-123"})
        assert login.status_code == 200, login.text
        headers = {"X-AssetGuard-Admin-Token": login.json()["access_token"]}

        workbook_response = await client.get("/admin/assets/export.xlsx", headers=headers)
        assert workbook_response.status_code == 200
        workbook = load_workbook(BytesIO(workbook_response.content), read_only=True, data_only=True)
        inventory_numbers = [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
        assert "SCOPE-VISIBLE" in inventory_numbers
        assert "SCOPE-HIDDEN" not in inventory_numbers

        pdf_response = await client.get("/admin/assets/export.pdf", headers=headers)
        assert pdf_response.status_code == 200
        with pdfplumber.open(BytesIO(pdf_response.content)) as exported_pdf:
            exported_text = "\n".join(page.extract_text() or "" for page in exported_pdf.pages)
        assert "SCOPE-VISIBLE" in exported_text
        assert "SCOPE-HIDDEN" not in exported_text

        visible_rooms = await client.get("/admin/vision/rooms", headers=headers)
        assert visible_rooms.status_code == 200
        assert [room["name"] for room in visible_rooms.json()] == ["Scope корпус A / 1 / 101"]
        assert visible_rooms.json()[0]["latest_scan"]["asset"] is None
        hidden_history = await client.get(f"/admin/vision/rooms/{hidden_room_id}/scans", headers=headers)
        hidden_detail = await client.get(f"/admin/vision/scans/{hidden_scan_id}", headers=headers)
        assert hidden_history.status_code == 404
        assert hidden_detail.status_code == 404

        hidden_inspections = await client.get(f"/admin/locations/rooms/{hidden_room_id}/inspections", headers=headers)
        assert hidden_inspections.status_code == 404
        viewer_write = await client.post(f"/admin/locations/rooms/{visible_room_id}/inspections", headers=headers, json={
            "items": [{"asset_id": str(visible_asset_id), "result": "PRESENT", "affected_quantity": 0}],
        })
        assert viewer_write.status_code == 404

        with factory() as session:
            grant = session.scalar(select(LocationAccessRecord).where(LocationAccessRecord.user_id == created_user_id))
            grant.permission = "EDITOR"
            session.commit()

        editor_write = await client.post(f"/admin/locations/rooms/{visible_room_id}/inspections", headers=headers, json={
            "comment": "Scoped editor inspection",
            "items": [{"asset_id": str(visible_asset_id), "result": "PRESENT", "affected_quantity": 0}],
        })
        assert editor_write.status_code == 201
        assert editor_write.json()["inspector_name"] == username
        visible_inspections = await client.get(f"/admin/locations/rooms/{visible_room_id}/inspections", headers=headers)
        assert visible_inspections.status_code == 200
        assert visible_inspections.json()[0]["comment"] == "Scoped editor inspection"

    with factory() as session:
        inspection_ids = select(RoomInspectionRecord.id).where(RoomInspectionRecord.room_id == visible_room_id)
        session.execute(text("ALTER TABLE room_inspection_items DISABLE TRIGGER trg_room_inspection_items_immutable"))
        session.execute(text("ALTER TABLE room_inspections DISABLE TRIGGER trg_room_inspections_immutable"))
        session.execute(delete(RoomInspectionItemRecord).where(RoomInspectionItemRecord.inspection_id.in_(inspection_ids)))
        session.execute(delete(RoomInspectionRecord).where(RoomInspectionRecord.room_id == visible_room_id))
        session.execute(text("ALTER TABLE room_inspection_items ENABLE TRIGGER trg_room_inspection_items_immutable"))
        session.execute(text("ALTER TABLE room_inspections ENABLE TRIGGER trg_room_inspections_immutable"))
        session.execute(text("ALTER TABLE asset_history_entries DISABLE TRIGGER trg_asset_history_immutable"))
        session.execute(delete(AssetHistoryEntryRecord).where(AssetHistoryEntryRecord.asset_id == visible_asset_id))
        session.execute(text("ALTER TABLE asset_history_entries ENABLE TRIGGER trg_asset_history_immutable"))
        session.execute(delete(AuthSessionRecord).where(AuthSessionRecord.user_id == created_user_id))
        session.execute(delete(LocationAccessRecord).where(LocationAccessRecord.user_id == created_user_id))
        session.execute(delete(UserRecord).where(UserRecord.id == created_user_id))
        session.execute(delete(VisionScanRecord).where(VisionScanRecord.room_id.in_([visible_vision_room.id, hidden_room_id])))
        session.execute(delete(VisionBaselineRecord).where(VisionBaselineRecord.room_id.in_([visible_vision_room.id, hidden_room_id])))
        session.execute(delete(VisionRoomRecord).where(VisionRoomRecord.organization_id == organization.id))
        session.execute(delete(AssetRecord).where(AssetRecord.inventory_number.in_(["SCOPE-VISIBLE", "SCOPE-HIDDEN"])))
        session.execute(delete(RoomRecord).where(RoomRecord.id.in_(room_ids)))
        session.execute(delete(FloorRecord).where(FloorRecord.id.in_(floor_ids)))
        session.execute(delete(BuildingRecord).where(BuildingRecord.id.in_(building_ids)))
        session.execute(delete(OrganizationRecord).where(OrganizationRecord.id == organization.id))
        session.commit()
