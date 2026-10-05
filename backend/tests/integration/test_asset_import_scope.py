from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from io import BytesIO

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from sqlalchemy import select

from assetguard.app import app
from assetguard.infrastructure.database import get_session_factory
from assetguard.modules.assets.models import AssetRecord, OrganizationRecord
from assetguard.modules.identity.auth import hash_password
from assetguard.modules.identity.models import UserRecord
from assetguard.infrastructure.config import get_settings


def import_file(kind: str, headers: list[str], rows: list[list[object]]) -> tuple[str, bytes, str]:
    output = BytesIO()
    if kind == "xlsx":
        workbook = Workbook()
        workbook.active.append(headers)
        for row in rows:
            workbook.active.append(row)
        workbook.save(output)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        table = Table([headers, *rows])
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
        SimpleDocTemplate(output).build([table])
        content_type = "application/pdf"
    return f"scope.{kind}", output.getvalue(), content_type


@pytest.mark.parametrize("kind", ["xlsx", "pdf"])
def test_import_row_errors_are_atomic_and_locations_remain_canonical(kind: str) -> None:
    async def exercise():
        now = datetime.now(UTC)
        with get_session_factory()() as session:
            own = OrganizationRecord(name="Canonical school", created_at=now)
            foreign = OrganizationRecord(name="Foreign school", created_at=now)
            session.add_all([own, foreign]); session.flush()
            session.add(UserRecord(username="canonical-admin", password_hash=hash_password("test-password"),
                                   role="ADMIN", organization_id=own.id, is_active=True, created_at=now))
            session.commit()
        columns = ["inventory_number", "name", "asset_type", "quantity", "unit", "building", "floor", "room"]
        valid = ["CAN-1", "Desks", "Furniture", 8, "шт", "Main", "1", "101"]
        bad = [valid, ["CAN-2", "", "Furniture", 2, "шт", "Main", "1", "101"],
               ["CAN-3", "Chairs", "Furniture", "1.5", "шт", "Main", "1", "101"]]
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as client:
            login = await client.post("/auth/login", json={"username": "canonical-admin", "password": "test-password"})
            headers = {"X-AssetGuard-Admin-Token": login.json()["access_token"]}
            path = f"/admin/assets/import.{kind}"
            for apply in (False, True):
                response = await client.post(path, headers=headers, params={"apply": str(apply).lower()},
                                             files={"file": import_file(kind, columns, bad)})
                assert response.status_code == 422
                detail = response.json()["detail"]
                assert detail["code"] == "IMPORT_ROWS_INVALID" and detail["error_count"] == 2
                assert [item["row"] for item in detail["errors"]] == [3, 4]
                assert all(item["page"] == (1 if kind == "pdf" else None) for item in detail["errors"])
                assert (await client.get("/admin/assets", headers=headers)).json() == []
            duplicate = await client.post(path, headers=headers, files={"file": import_file(kind, columns, [valid, valid])})
            assert duplicate.status_code == 422
            assert duplicate.json()["detail"]["errors"][0]["row"] == 3
            good_file = import_file(kind, columns, [valid])
            preview = await client.post(path, headers=headers, files={"file": good_file})
            assert preview.status_code == 200
            assert preview.json()["scope"]["organization_name"] == "Canonical school"
            assert preview.json()["items"][0]["source_row"] == 2
            assert (await client.get("/admin/locations/tree", headers=headers)).json() == []
            for _ in range(2):
                assert (await client.post(path, params={"apply": "true"}, headers=headers, files={"file": good_file})).status_code == 200
            assets = (await client.get("/admin/assets", headers=headers)).json()
            assert len(assets) == 1 and assets[0]["room_id"]
            room_id = assets[0]["room_id"]
            tree = (await client.get("/admin/locations/tree", headers=headers)).json()
            assert len(tree) == 1 and tree[0]["organization_id"] == preview.json()["scope"]["organization_id"]
            workspace = (await client.get(f"/admin/locations/rooms/{room_id}/workspace", headers=headers)).json()
            assert workspace["inventory"]["assets"][0]["quantity"] == 8
            old_format = import_file(kind, ["inventory_number", "name", "asset_type"], [["CAN-1", "Updated desks", "Furniture"]])
            assert (await client.post(path, params={"apply": "true"}, headers=headers, files={"file": old_format})).status_code == 200
            updated = (await client.get("/admin/assets", headers=headers)).json()[0]
            assert updated["room_id"] == room_id and updated["room"] == "101" and updated["quantity"] == 8
            foreign_file = import_file(kind, ["inventory_number", "name", "asset_type", "organization"],
                                       [["FORBIDDEN", "Desk", "Furniture", "Foreign school"]])
            assert (await client.post(path, params={"apply": "true"}, headers=headers, files={"file": foreign_file})).status_code == 403
            assert len((await client.get("/admin/assets", headers=headers)).json()) == 1
    asyncio.run(exercise())


@pytest.mark.parametrize("kind", ["xlsx", "pdf"])
def test_import_uses_authenticated_organization_for_preview_create_and_update(kind: str) -> None:
    async def exercise() -> None:
        now = datetime.now(UTC)
        with get_session_factory()() as session:
            own = OrganizationRecord(name="Import school", created_at=now)
            default = OrganizationRecord(name="Default Organization", created_at=now)
            session.add_all([own, default])
            session.flush()
            session.add(UserRecord(username="import-admin", password_hash=hash_password("import-test-password"), role="ADMIN", organization_id=own.id, is_active=True, created_at=now))
            for organization, name in ((own, "Own original"), (default, "Foreign original")):
                session.add(AssetRecord(organization_id=organization.id, inventory_number="SAME", name=name, asset_type="Other", status="ACTIVE", created_at=now, updated_at=now))
            session.commit()
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
            login = await client.post("/auth/login", json={"username": "import-admin", "password": "import-test-password"})
            assert login.status_code == 200, login.text
            auth = {"X-AssetGuard-Admin-Token": login.json()["access_token"]}
            file = import_file(kind, ["inventory_number", "name", "asset_type"], [["SAME", "Own updated", "Other"], ["NEW", "Own created", "Other"]])
            preview = await client.post(f"/admin/assets/import.{kind}", headers=auth, files={"file": file})
            assert preview.status_code == 200, preview.text
            assert (preview.json()["creates"], preview.json()["updates"]) == (1, 1)
            assert {item["organization"] for item in preview.json()["items"]} == {own.name}
            applied = await client.post(f"/admin/assets/import.{kind}?apply=true", headers=auth, files={"file": file})
            assert applied.status_code == 200, applied.text
            with get_session_factory()() as session:
                assets = list(session.scalars(select(AssetRecord)))
                assert {(a.organization_id, a.inventory_number, a.name) for a in assets} == {
                    (own.id, "SAME", "Own updated"), (own.id, "NEW", "Own created"), (default.id, "SAME", "Foreign original"),
                }
            mixed = import_file(kind, ["inventory_number", "name", "asset_type", "organization"], [["NEW", "Forbidden update", "Other", own.name], ["FOREIGN", "Forbidden create", "Other", default.name]])
            for apply in (False, True):
                response = await client.post(f"/admin/assets/import.{kind}?apply={str(apply).lower()}", headers=auth, files={"file": mixed})
                assert response.status_code == 403, response.text
            with get_session_factory()() as session:
                assert session.scalar(select(AssetRecord).where(AssetRecord.inventory_number == "NEW")).name == "Own created"
                assert session.scalar(select(AssetRecord).where(AssetRecord.inventory_number == "FOREIGN")) is None
    asyncio.run(exercise())


def test_grouped_import_create_update_repeat_and_export_preserve_accounting() -> None:
    async def exercise() -> None:
        transport = httpx.ASGITransport(app=app)
        auth = {"X-AssetGuard-Admin-Token": get_settings().admin_shared_secret}
        async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
            columns = ["inventory_number", "name", "asset_type", "quantity", "unit", "tracking_mode"]
            file = import_file("xlsx", columns, [["GROUP", "Desks", "Furniture", "12,000", "шт", "GROUPED"]])
            preview = await client.post("/admin/assets/import.xlsx", headers=auth, files={"file": file})
            assert preview.status_code == 200, preview.text
            item = preview.json()["items"][0]
            assert (item["quantity"], item["unit"], item["tracking_mode"]) == (12, "шт", "GROUPED")
            for _ in range(2):
                response = await client.post("/admin/assets/import.xlsx?apply=true", headers=auth, files={"file": file})
                assert response.status_code == 200, response.text
            changed = import_file("xlsx", columns, [["GROUP", "Updated desks", "Furniture", 10, "компл", "GROUPED"]])
            assert (await client.post("/admin/assets/import.xlsx?apply=true", headers=auth, files={"file": changed})).status_code == 200
            legacy = import_file("xlsx", columns[:3], [["GROUP", "Legacy name", "Furniture"]])
            assert (await client.post("/admin/assets/import.xlsx?apply=true", headers=auth, files={"file": legacy})).status_code == 200
            assets = (await client.get("/admin/assets", headers=auth)).json()
            assert len(assets) == 1
            assert (assets[0]["quantity"], assets[0]["unit"], assets[0]["tracking_mode"]) == (10, "компл", "GROUPED")
            exported = await client.get("/admin/assets/export.xlsx", headers=auth)
            workbook = load_workbook(BytesIO(exported.content), data_only=True)
            header, row = list(workbook.active.values)
            values = dict(zip(header, row))
            assert (values["quantity"], values["unit"], values["tracking_mode"]) == (10, "компл", "GROUPED")
            assert (await client.post("/admin/assets/import.xlsx?apply=true", headers=auth, files={"file": ("export.xlsx", exported.content)})).status_code == 200
    asyncio.run(exercise())
