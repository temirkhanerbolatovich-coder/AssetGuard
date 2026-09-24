from __future__ import annotations

import secrets
from io import BytesIO
from datetime import UTC, datetime
from typing import Annotated, Literal
from urllib.parse import urlsplit
from uuid import UUID

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile, status
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook, load_workbook
import segno
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from assetguard.infrastructure.config import get_settings
from assetguard.infrastructure.database import get_session
from assetguard.modules.assets.models import AssetRecord, OrganizationRecord
from assetguard.modules.baselines.models import BaselineRecord
from assetguard.modules.changes.models import ChangeEventRecord
from assetguard.modules.history.service import append_asset_history
from assetguard.modules.endpoints.service import evaluate_last_seen
from assetguard.modules.incidents.models import AssetHistoryEntryRecord, IncidentRecord
from assetguard.modules.inventory.models import RawInventoryRecord
from assetguard.modules.identity.auth import AuthPrincipal, session_principal
from assetguard.modules.snapshots.models import (
    ComponentObservationRecord, EndpointIdentifierRecord,
    HardwareSnapshotRecord, ManagedEndpointRecord,
)

router = APIRouter(prefix="/admin", tags=["admin"])


def require_admin(
    token: Annotated[str | None, Header(alias="X-AssetGuard-Admin-Token")] = None,
    session: Session = Depends(get_session),
) -> AuthPrincipal:
    settings = get_settings()
    valid = [settings.admin_shared_secret, settings.previous_admin_shared_secret]
    shared = bool(token and any(candidate and secrets.compare_digest(token, candidate) for candidate in valid))
    if shared:
        return AuthPrincipal(username="bootstrap-admin", role="ADMIN", session_id=None)
    principal = session_principal(session, token) if token else None
    if not principal or principal.role != "ADMIN":
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid administrator credentials.")
    return principal


def require_viewer(
    token: Annotated[str | None, Header(alias="X-AssetGuard-Admin-Token")] = None,
    session: Session = Depends(get_session),
) -> AuthPrincipal:
    settings = get_settings()
    valid = [settings.admin_shared_secret, settings.previous_admin_shared_secret, settings.viewer_shared_secret]
    if token and any(candidate and secrets.compare_digest(token, candidate) for candidate in valid):
        role = "VIEWER" if settings.viewer_shared_secret and secrets.compare_digest(token, settings.viewer_shared_secret) else "ADMIN"
        return AuthPrincipal(username=f"shared-{role.lower()}", role=role, session_id=None)
    principal = session_principal(session, token) if token else None
    if not principal or principal.role not in {"ADMIN", "VIEWER"}:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid AssetGuard credentials.")
    return principal


class AssetCreate(BaseModel):
    inventory_number: str = Field(min_length=1, max_length=128)
    name: str = Field(min_length=1, max_length=255)
    asset_type: Literal["Desktop", "Laptop", "Other"]
    status: str = Field(default="ACTIVE", max_length=32)
    building: str | None = Field(default=None, max_length=255)
    floor: str | None = Field(default=None, max_length=64)
    room: str | None = Field(default=None, max_length=255)
    notes: str | None = None


class AssetUpdate(BaseModel):
    inventory_number: str | None = Field(default=None, min_length=1, max_length=128)
    name: str | None = Field(default=None, min_length=1, max_length=255)
    asset_type: Literal["Desktop", "Laptop", "Other"] | None = None
    status: str | None = Field(default=None, max_length=32)
    building: str | None = Field(default=None, max_length=255)
    floor: str | None = Field(default=None, max_length=64)
    room: str | None = Field(default=None, max_length=255)
    notes: str | None = None


def _asset_view(asset: AssetRecord, endpoint_id: UUID | None, organization: str | None = None) -> dict:
    return {
        "id": str(asset.id), "inventory_number": asset.inventory_number,
        "name": asset.name, "asset_type": asset.asset_type, "status": asset.status,
        "building": asset.building, "floor": asset.floor, "room": asset.room,
        "notes": asset.notes, "organization": organization,
        "endpoint_id": str(endpoint_id) if endpoint_id else None,
    }


def _organization_for_name(session: Session, name: str) -> OrganizationRecord:
    organization = session.scalar(select(OrganizationRecord).where(OrganizationRecord.name == name))
    if organization:
        return organization
    organization = OrganizationRecord(name=name, created_at=datetime.now(UTC))
    session.add(organization)
    session.flush()
    return organization


def _latest_snapshot(session: Session, endpoint_id: UUID) -> HardwareSnapshotRecord | None:
    return session.scalar(select(HardwareSnapshotRecord).where(
        HardwareSnapshotRecord.managed_endpoint_id == endpoint_id,
    ).order_by(HardwareSnapshotRecord.captured_at.desc()))


def _latest_observed_snapshot(session: Session, endpoint_id: UUID) -> HardwareSnapshotRecord | None:
    snapshots = session.scalars(select(HardwareSnapshotRecord).where(
        HardwareSnapshotRecord.managed_endpoint_id == endpoint_id,
    ).order_by(HardwareSnapshotRecord.captured_at.desc()).limit(50))
    return next((snapshot for snapshot in snapshots if session.scalar(
        select(func.count()).select_from(ComponentObservationRecord).where(
            ComponentObservationRecord.hardware_snapshot_id == snapshot.id,
        )
    )), None)


def _latest_components(session: Session, endpoint_id: UUID) -> list[ComponentObservationRecord]:
    result: list[ComponentObservationRecord] = []
    observed_types: set[str] = set()
    snapshots = session.scalars(select(HardwareSnapshotRecord).where(
        HardwareSnapshotRecord.managed_endpoint_id == endpoint_id,
    ).order_by(HardwareSnapshotRecord.captured_at.desc()).limit(50))
    for snapshot in snapshots:
        observations = list(session.scalars(select(ComponentObservationRecord).where(
            ComponentObservationRecord.hardware_snapshot_id == snapshot.id,
        )))
        for component_type in {item.component_type for item in observations} - observed_types:
            result.extend(item for item in observations if item.component_type == component_type)
            observed_types.add(component_type)
    return result


def _ram_capacity_bytes(value: int | None) -> int | None:
    """Accept both GLPI's legacy MiB values and byte-based inventory payloads."""
    if value is None:
        return None
    return value * 1024 * 1024 if value < 1024 * 1024 else value


def _component_view(item: ComponentObservationRecord) -> dict:
    return {
        "type": item.component_type, "model": item.model, "serial": item.serial_number,
        "manufacturer": item.manufacturer, "part_number": item.part_number,
        "capacity": _ram_capacity_bytes(item.capacity) if item.component_type == "RAM" else item.capacity,
        "slot": item.slot, "confidence": item.confidence,
        "raw_data": item.raw_data,
    }


def _endpoint_summary(session: Session, endpoint: ManagedEndpointRecord) -> dict:
    snapshot = _latest_snapshot(session, endpoint.id)
    observations = _latest_components(session, endpoint.id)
    open_changes = session.scalar(select(func.count()).select_from(ChangeEventRecord).where(
        ChangeEventRecord.managed_endpoint_id == endpoint.id,
        ChangeEventRecord.status == "OPEN",
    )) or 0
    open_incidents = session.scalar(select(func.count()).select_from(IncidentRecord).where(
        IncidentRecord.managed_endpoint_id == endpoint.id,
        IncidentRecord.status.in_(("OPEN", "UNDER_REVIEW")),
    )) or 0
    ram = [item for item in observations if item.component_type == "RAM"]
    storage = [item for item in observations if item.component_type == "STORAGE"]
    cpu = next((item for item in observations if item.component_type == "CPU"), None)
    return {
        "id": str(endpoint.id), "asset_id": str(endpoint.asset_id) if endpoint.asset_id else None,
        "source": endpoint.source, "source_agent_id": endpoint.source_agent_id,
        "hostname": endpoint.hostname, "last_seen_at": endpoint.last_seen_at,
        "status": endpoint.status, "open_changes": open_changes,
        "open_incidents": open_incidents,
        "current_snapshot": None if not snapshot else {
            "id": str(snapshot.id), "captured_at": snapshot.captured_at,
            "type": snapshot.snapshot_type, "completeness": snapshot.completeness,
        },
        "hardware_summary": {
            "ram_bytes": sum(_ram_capacity_bytes(item.capacity) or 0 for item in ram),
            "ram_modules": len(ram), "storage_devices": len(storage),
            "cpu": cpu.model if cpu else None,
        },
    }


@router.get("/assets", dependencies=[Depends(require_viewer)])
def list_assets(session: Annotated[Session, Depends(get_session)]):
    result = []
    for asset in session.scalars(select(AssetRecord).order_by(AssetRecord.inventory_number)):
        endpoint = session.scalar(select(ManagedEndpointRecord).where(ManagedEndpointRecord.asset_id == asset.id))
        organization = session.get(OrganizationRecord, asset.organization_id)
        view = _asset_view(asset, endpoint.id if endpoint else None, organization.name if organization else None)
        view["endpoint"] = _endpoint_summary(session, endpoint) if endpoint else None
        result.append(view)
    return result


@router.get("/assets/export.xlsx", dependencies=[Depends(require_viewer)])
def export_assets_xlsx(session: Annotated[Session, Depends(get_session)]):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Assets"
    headers = ["inventory_number", "name", "asset_type", "status", "organization", "building", "floor", "room", "notes"]
    sheet.append(headers)
    for asset in session.scalars(select(AssetRecord).order_by(AssetRecord.inventory_number)):
        organization = session.get(OrganizationRecord, asset.organization_id)
        sheet.append([
            asset.inventory_number, asset.name, asset.asset_type, asset.status,
            organization.name if organization else "", asset.building or "", asset.floor or "",
            asset.room or "", asset.notes or "",
        ])
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="assetguard-assets.xlsx"'},
    )


IMPORT_COLUMN_ALIASES = {
    "inventory_number": {"inventory_number", "inventory number", "инвентарный номер", "инв. номер", "инв номер", "инвентарный№"},
    "name": {"name", "название", "наименование", "оборудование", "устройство"},
    "asset_type": {"asset_type", "asset type", "тип", "тип оборудования", "вид оборудования"},
    "status": {"status", "статус", "состояние"},
    "organization": {"organization", "организация", "учреждение", "школа"},
    "building": {"building", "корпус", "здание"},
    "floor": {"floor", "этаж"},
    "room": {"room", "кабинет", "аудитория", "помещение"},
    "notes": {"notes", "примечание", "комментарий"},
}


def _normalize_import_header(value: object) -> str:
    return " ".join(str(value).strip().lower().replace("ё", "е").split())


def _canonical_import_columns(header: tuple[object, ...]) -> dict[str, int]:
    raw = {_normalize_import_header(value): index for index, value in enumerate(header) if value is not None}
    return {
        field: next((raw[alias] for alias in aliases if alias in raw), None)
        for field, aliases in IMPORT_COLUMN_ALIASES.items()
    }


def _asset_type_from_import(value: str | None) -> str | None:
    if not value:
        return value
    normalized = value.lower().replace("ё", "е").strip()
    if normalized in {"desktop", "пк", "компьютер", "стационарный", "стационарный компьютер"}:
        return "Desktop"
    if normalized in {"laptop", "ноутбук"}:
        return "Laptop"
    return "Other" if normalized not in {"other", "прочее", "другое"} else "Other"


def _import_row_values(row: dict[str, object], row_number: int) -> dict[str, str | None]:
    required = ("inventory_number", "name", "asset_type")
    result = {key: (str(row.get(key)).strip() if row.get(key) is not None else None) for key in (
        "inventory_number", "name", "asset_type", "status", "organization", "building", "floor", "room", "notes",
    )}
    if any(not result[key] for key in required):
        raise ValueError(f"Row {row_number}: inventory_number, name and asset_type are required.")
    result["asset_type"] = _asset_type_from_import(result["asset_type"])
    if len(result["inventory_number"] or "") > 128 or len(result["name"] or "") > 255:
        raise ValueError(f"Row {row_number}: inventory_number or name is too long.")
    return result


@router.post("/assets/import.xlsx", dependencies=[Depends(require_admin)])
def import_assets_xlsx(
    file: Annotated[UploadFile, File()],
    session: Annotated[Session, Depends(get_session)],
    apply: bool = Query(default=False),
):
    if file.content_type not in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    } and not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "Upload an .xlsx file.")
    try:
        workbook = load_workbook(BytesIO(file.file.read(5 * 1024 * 1024 + 1)), read_only=True, data_only=True)
    except Exception as error:
        raise HTTPException(422, "The uploaded file is not a valid .xlsx workbook.") from error
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise HTTPException(422, "The workbook is empty.")
    columns = _canonical_import_columns(header)
    required_columns = {"inventory_number", "name", "asset_type"}
    if any(columns[column] is None for column in required_columns):
        raise HTTPException(422, "Required columns: inventory_number, name, asset_type (or Russian equivalents).")
    parsed: list[dict[str, str | None]] = []
    try:
        for row_number, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            parsed.append(_import_row_values({name: values[index] if index is not None and index < len(values) else None for name, index in columns.items()}, row_number))
    except ValueError as error:
        raise HTTPException(422, str(error)) from error
    existing = {(organization.name, asset.inventory_number): asset for asset in session.scalars(select(AssetRecord).join(OrganizationRecord)) for organization in [session.get(OrganizationRecord, asset.organization_id)] if organization}
    creates = updates = 0
    for item in parsed:
        organization_name = item["organization"] or "Default Organization"
        if (organization_name, item["inventory_number"]) in existing:
            updates += 1
        else:
            creates += 1
    if not apply:
        return {"rows": len(parsed), "creates": creates, "updates": updates, "applied": False}
    now = datetime.now(UTC)
    for item in parsed:
        organization = _organization_for_name(session, item["organization"] or "Default Organization")
        asset = existing.get((organization.name, item["inventory_number"]))
        if asset is None:
            asset = AssetRecord(
                organization_id=organization.id, inventory_number=item["inventory_number"] or "",
                name=item["name"] or "", asset_type=item["asset_type"] or "Other",
                status=item["status"] or "ACTIVE", building=item["building"], floor=item["floor"],
                room=item["room"], notes=item["notes"], created_at=now, updated_at=now,
            )
            session.add(asset)
            session.flush()
            append_asset_history(session, asset_id=asset.id, event_type="ASSET_CREATED", related_entity_type="Asset", related_entity_id=asset.id, message="Asset imported from workbook.", metadata={"inventory_number": asset.inventory_number})
        else:
            for key in ("name", "asset_type", "building", "floor", "room", "notes"):
                setattr(asset, key, item[key])
            if item["status"]:
                asset.status = item["status"]
            asset.updated_at = now
            append_asset_history(session, asset_id=asset.id, event_type="ASSET_UPDATED", related_entity_type="Asset", related_entity_id=asset.id, message="Asset imported from workbook.", metadata={"source": "xlsx"})
    session.commit()
    return {"rows": len(parsed), "creates": creates, "updates": updates, "applied": True}


@router.get("/assets/{asset_id}/qr.svg", dependencies=[Depends(require_viewer)])
def asset_qr_svg(
    asset_id: UUID,
    session: Annotated[Session, Depends(get_session)],
    public_url: str | None = Query(default=None, max_length=2048),
):
    asset = session.get(AssetRecord, asset_id)
    if not asset:
        raise HTTPException(404, "Asset was not found.")
    base_url = get_settings().public_url
    if public_url:
        parsed = urlsplit(public_url)
        if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:
            raise HTTPException(422, "QR public_url must be an HTTPS origin without credentials, query or fragment.")
        base_url = f"https://{parsed.netloc}{parsed.path.rstrip('/')}"
    payload = f"{base_url}/#asset={asset.id}"
    qr = segno.make(payload, error="m")
    output = BytesIO()
    qr.save(output, kind="svg", scale=4, border=2, title=f"AssetGuard {asset.inventory_number}")
    return Response(output.getvalue(), media_type="image/svg+xml")


@router.post("/assets", dependencies=[Depends(require_admin)], status_code=status.HTTP_201_CREATED)
def create_asset(body: AssetCreate, session: Annotated[Session, Depends(get_session)]):
    organization = session.scalar(select(OrganizationRecord).order_by(OrganizationRecord.created_at))
    if not organization:
        organization = OrganizationRecord(name="Default Organization", created_at=datetime.now(UTC))
        session.add(organization)
        session.flush()
    if session.scalar(select(AssetRecord).where(
        AssetRecord.organization_id == organization.id,
        AssetRecord.inventory_number == body.inventory_number,
    )):
        raise HTTPException(status.HTTP_409_CONFLICT, "Inventory number already exists.")
    now = datetime.now(UTC)
    asset = AssetRecord(
        organization_id=organization.id, inventory_number=body.inventory_number,
        name=body.name, asset_type=body.asset_type, status=body.status,
        building=body.building, floor=body.floor, room=body.room, notes=body.notes,
        created_at=now, updated_at=now,
    )
    session.add(asset)
    session.flush()
    append_asset_history(
        session, asset_id=asset.id, event_type="ASSET_CREATED",
        related_entity_type="Asset", related_entity_id=asset.id,
        message="Asset created.", metadata={"inventory_number": asset.inventory_number},
    )
    session.commit()
    session.refresh(asset)
    return _asset_view(asset, None, organization.name)


@router.patch("/assets/{asset_id}", dependencies=[Depends(require_admin)])
def update_asset(asset_id: UUID, body: AssetUpdate, session: Annotated[Session, Depends(get_session)]):
    asset = session.get(AssetRecord, asset_id)
    if not asset:
        raise HTTPException(404, "Asset was not found.")
    changes = body.model_dump(exclude_unset=True)
    if "inventory_number" in changes:
        duplicate = session.scalar(select(AssetRecord).where(
            AssetRecord.organization_id == asset.organization_id,
            AssetRecord.inventory_number == changes["inventory_number"],
            AssetRecord.id != asset.id,
        ))
        if duplicate:
            raise HTTPException(409, "Inventory number already exists.")
    for field, value in changes.items():
        setattr(asset, field, value)
    asset.updated_at = datetime.now(UTC)
    append_asset_history(
        session, asset_id=asset.id, event_type="ASSET_UPDATED",
        related_entity_type="Asset", related_entity_id=asset.id,
        message="Asset details updated.", metadata={"fields": sorted(changes)},
    )
    session.commit()
    endpoint_id = session.scalar(select(ManagedEndpointRecord.id).where(ManagedEndpointRecord.asset_id == asset.id))
    organization = session.get(OrganizationRecord, asset.organization_id)
    return _asset_view(asset, endpoint_id, organization.name if organization else None)


@router.get("/endpoints", dependencies=[Depends(require_viewer)])
def list_endpoints(session: Annotated[Session, Depends(get_session)]):
    result = []
    for endpoint in session.scalars(select(ManagedEndpointRecord).order_by(ManagedEndpointRecord.last_seen_at.desc())):
        item = _endpoint_summary(session, endpoint)
        asset = session.get(AssetRecord, endpoint.asset_id) if endpoint.asset_id else None
        organization = session.get(OrganizationRecord, asset.organization_id) if asset else None
        item["asset"] = None if not asset else _asset_view(asset, endpoint.id, organization.name if organization else None)
        result.append(item)
    return result


@router.post("/maintenance/evaluate-endpoints", dependencies=[Depends(require_admin)])
def evaluate_endpoints(session: Annotated[Session, Depends(get_session)]):
    changed = evaluate_last_seen(session, get_settings().endpoint_stale_after_hours)
    return {"updated": changed, "stale_after_hours": get_settings().endpoint_stale_after_hours}


@router.get("/endpoints/{endpoint_id}", dependencies=[Depends(require_viewer)])
def endpoint_detail(endpoint_id: UUID, session: Annotated[Session, Depends(get_session)]):
    endpoint = session.get(ManagedEndpointRecord, endpoint_id)
    if not endpoint:
        raise HTTPException(404, "Endpoint was not found.")
    identifiers = list(session.scalars(select(EndpointIdentifierRecord).where(
        EndpointIdentifierRecord.managed_endpoint_id == endpoint.id,
    ).order_by(EndpointIdentifierRecord.identifier_type)))
    snapshots = list(session.scalars(select(HardwareSnapshotRecord).where(
        HardwareSnapshotRecord.managed_endpoint_id == endpoint.id,
    ).order_by(HardwareSnapshotRecord.captured_at.desc())))
    return {
        "id": str(endpoint.id), "asset_id": str(endpoint.asset_id) if endpoint.asset_id else None,
        "source": endpoint.source, "source_agent_id": endpoint.source_agent_id,
        "hostname": endpoint.hostname, "last_seen_at": endpoint.last_seen_at,
        "status": endpoint.status,
        "identifiers": [{
            "type": item.identifier_type, "value": item.raw_value,
            "confidence": item.confidence, "first_seen_at": item.first_seen_at,
            "last_seen_at": item.last_seen_at, "active": item.is_active,
        } for item in identifiers],
        "snapshot_count": len(snapshots),
        "current_snapshot_id": str(snapshots[0].id) if snapshots else None,
    }


@router.post("/endpoints/{endpoint_id}/asset/{asset_id}", dependencies=[Depends(require_admin)])
def link_endpoint(endpoint_id: UUID, asset_id: UUID, session: Annotated[Session, Depends(get_session)]):
    endpoint = session.get(ManagedEndpointRecord, endpoint_id)
    asset = session.get(AssetRecord, asset_id)
    if not endpoint or not asset:
        raise HTTPException(404, "Asset or endpoint was not found.")
    if endpoint.asset_id and endpoint.asset_id != asset.id:
        append_asset_history(
            session, asset_id=endpoint.asset_id, endpoint_id=endpoint.id,
            event_type="ENDPOINT_UNLINKED", related_entity_type="ManagedEndpoint",
            related_entity_id=endpoint.id, message="Endpoint unlinked from asset.",
        )
    for prior in session.scalars(select(ManagedEndpointRecord).where(
        ManagedEndpointRecord.asset_id == asset.id, ManagedEndpointRecord.id != endpoint.id,
    )):
        append_asset_history(
            session, asset_id=asset.id, endpoint_id=prior.id,
            event_type="ENDPOINT_UNLINKED", related_entity_type="ManagedEndpoint",
            related_entity_id=prior.id, message="Previous endpoint unlinked from asset.",
        )
        prior.asset_id = None
        prior.updated_at = datetime.now(UTC)
    endpoint.asset_id = asset.id
    endpoint.updated_at = datetime.now(UTC)
    append_asset_history(
        session, asset_id=asset.id, endpoint_id=endpoint.id,
        event_type="ENDPOINT_LINKED", related_entity_type="ManagedEndpoint",
        related_entity_id=endpoint.id, message="Endpoint linked to asset.",
        metadata={"hostname": endpoint.hostname},
    )
    session.commit()
    return {"status": "linked"}


def _components(session: Session, snapshot: HardwareSnapshotRecord | None) -> list[dict]:
    if not snapshot:
        return []
    return [_component_view(item) for item in session.scalars(select(ComponentObservationRecord).where(
        ComponentObservationRecord.hardware_snapshot_id == snapshot.id,
    ))]


def _system_sections(
    session: Session, endpoint: ManagedEndpointRecord, snapshot: HardwareSnapshotRecord | None,
) -> tuple[dict, dict | None]:
    if not snapshot:
        return {}, None
    latest_raw = session.get(RawInventoryRecord, snapshot.raw_inventory_id)
    if not latest_raw:
        return {}, None
    inventories = list(session.scalars(select(RawInventoryRecord).where(
        RawInventoryRecord.managed_endpoint_id == endpoint.id,
        RawInventoryRecord.processing_status == "PROCESSED",
    ).order_by(RawInventoryRecord.received_at.desc()).limit(50)))
    merged_objects = {"hardware": {}, "bios": {}, "operatingsystem": {}, "assetguard_network": {}}
    latest_lists = {"drives": [], "controllers": []}
    for inventory in inventories:
        content = inventory.payload.get("content") if isinstance(inventory.payload, dict) else None
        if not isinstance(content, dict):
            continue
        for section in merged_objects:
            value = content.get(section)
            if isinstance(value, dict):
                for key, item in value.items():
                    if key not in merged_objects[section] and item not in (None, ""):
                        merged_objects[section][key] = item
        for section in latest_lists:
            value = content.get(section)
            if not latest_lists[section] and isinstance(value, list) and value:
                latest_lists[section] = value
    hardware = merged_objects["hardware"]
    # Windows product keys/owner fields remain evidence-only and are never exposed to the dashboard.
    safe_hardware = {key: hardware.get(key) for key in (
        "name", "uuid", "chassis_type", "memory", "workgroup", "winlang", "vmsystem",
    ) if hardware.get(key) not in (None, "")}
    return {
        "hardware": safe_hardware,
        "bios": merged_objects["bios"],
        "operating_system": merged_objects["operatingsystem"],
        "network_quality": merged_objects["assetguard_network"],
        "drives": latest_lists["drives"],
        "controllers": latest_lists["controllers"],
    }, {
        "id": str(latest_raw.id), "received_at": latest_raw.received_at, "source": latest_raw.source,
        "source_version": latest_raw.source_version, "type": latest_raw.inventory_type,
        "processing_status": latest_raw.processing_status,
    }


@router.get("/assets/{asset_id}", dependencies=[Depends(require_viewer)])
def asset_detail(asset_id: UUID, session: Annotated[Session, Depends(get_session)]):
    asset = session.get(AssetRecord, asset_id)
    if not asset:
        raise HTTPException(404, "Asset was not found.")
    endpoint = session.scalar(select(ManagedEndpointRecord).where(ManagedEndpointRecord.asset_id == asset.id))
    organization = session.get(OrganizationRecord, asset.organization_id)
    result = _asset_view(asset, endpoint.id if endpoint else None, organization.name if organization else None)
    result.update({
        "endpoint": None, "current_snapshot_id": None, "baseline_snapshot_id": None,
        "current_hardware": [], "baseline_hardware": [], "changes": [], "incidents": [],
        "system": {}, "latest_inventory": None,
        "history": [{
            "id": str(item.id), "type": item.event_type, "occurred_at": item.occurred_at,
            "message": item.message, "metadata": item.metadata_json,
        } for item in session.scalars(select(AssetHistoryEntryRecord).where(
            AssetHistoryEntryRecord.asset_id == asset.id,
        ).order_by(AssetHistoryEntryRecord.occurred_at.desc()))],
    })
    if not endpoint:
        return result
    snapshots = list(session.scalars(select(HardwareSnapshotRecord).where(
        HardwareSnapshotRecord.managed_endpoint_id == endpoint.id,
    ).order_by(HardwareSnapshotRecord.captured_at.desc())))
    current = snapshots[0] if snapshots else None
    baseline = session.scalar(select(BaselineRecord).where(
        BaselineRecord.managed_endpoint_id == endpoint.id, BaselineRecord.status == "ACTIVE",
    ))
    baseline_snapshot = session.get(HardwareSnapshotRecord, baseline.hardware_snapshot_id) if baseline else None
    identifiers = list(session.scalars(select(EndpointIdentifierRecord).where(
        EndpointIdentifierRecord.managed_endpoint_id == endpoint.id,
        EndpointIdentifierRecord.is_active.is_(True),
    ).order_by(EndpointIdentifierRecord.identifier_type)))
    observed_snapshot = _latest_observed_snapshot(session, endpoint.id)
    system, latest_inventory = _system_sections(session, endpoint, current)
    result.update({
        "endpoint": {
            **_endpoint_summary(session, endpoint),
            "identifiers": [{
                "type": item.identifier_type, "value": item.raw_value,
                "confidence": item.confidence,
            } for item in identifiers],
        },
        "current_snapshot_id": str(current.id) if current else None,
        "recommended_baseline_snapshot_id": str(observed_snapshot.id) if observed_snapshot else None,
        "baseline_snapshot_id": str(baseline_snapshot.id) if baseline_snapshot else None,
        "current_snapshot": None if not current else {
            "id": str(current.id), "captured_at": current.captured_at,
            "type": current.snapshot_type, "completeness": current.completeness,
        },
        "baseline": None if not baseline else {
            "id": str(baseline.id), "snapshot_id": str(baseline.hardware_snapshot_id),
            "accepted_at": baseline.accepted_at, "reason": baseline.reason,
        },
        "current_hardware": [_component_view(item) for item in _latest_components(session, endpoint.id)],
        "baseline_hardware": _components(session, baseline_snapshot),
        "system": system, "latest_inventory": latest_inventory,
        "changes": [{
            "id": str(item.id), "type": item.event_type,
            "component_type": item.component_type, "confidence": item.confidence,
            "severity": item.severity, "status": item.status,
            "evidence": item.evidence, "detected_at": item.detected_at,
        } for item in session.scalars(select(ChangeEventRecord).where(
            ChangeEventRecord.managed_endpoint_id == endpoint.id,
        ).order_by(ChangeEventRecord.detected_at.desc()))],
        "incidents": [{
            "id": str(item.id), "title": item.title, "status": item.status,
            "severity": item.severity, "created_at": item.created_at,
        } for item in session.scalars(select(IncidentRecord).where(
            IncidentRecord.managed_endpoint_id == endpoint.id,
        ).order_by(IncidentRecord.created_at.desc()))],
    })
    return result
